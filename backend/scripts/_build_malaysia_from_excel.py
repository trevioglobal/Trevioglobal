"""
Parse Malaysia Contracted rates + KTH RATE SHEET Excels → JSON for TS importer.
Uses rounded USD (matches printed rate sheet / image) and USD×84 → INR.
"""
from __future__ import annotations

import json
import math
import re
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(r"c:\Users\ADMIN\Desktop\projects\travelpartner-trevio")
HOTEL_XLSX = ROOT / "docs" / "Malaysia Contracted rates .xlsx"
KTH_XLSX = ROOT / "docs" / "KTH RATE SHEET 2026 - TREVIO.xlsx"
OUT = ROOT / "scripts" / "malaysia-from-excel.json"

USD_TO_INR = 84.0
MYR_TO_INR = 19.0

HOTEL_TERMS_VALID_TO = "2026-09-30"
KTH_VALID_FROM = "2026-01-01"
KTH_VALID_TO = "2026-12-31"


def money_usd(v) -> float:
    return float(v)


def round_usd(v) -> int:
    """Printed sheet / image uses whole USD; Excel stores fractional ROE values."""
    return int(round(float(v)))


def inr_from_usd(usd: int) -> int:
    return int(round(usd * USD_TO_INR))


def inr_from_myr(myr) -> int | None:
    if myr is None or myr == "":
        return None
    try:
        return int(round(float(myr) * MYR_TO_INR))
    except Exception:
        return None


def parse_hotel_header(text: str) -> tuple[str, int] | None:
    t = re.sub(r"\s+", " ", str(text or "").strip())
    if not t or t.lower().startswith("terms"):
        return None
    # Dash seasonal headers are hotel headers too
    m = re.match(r"^(.+?)\s+(\d)\*?\s*(?:Star Premier)?\s*(?:-\s*.+)?$", t, re.I)
    if not m:
        # e.g. "Trigo Kuala Lumpur ( 3 Star Premier )"
        m2 = re.match(r"^(.+?)\s*\(\s*(\d)\s*Star", t, re.I)
        if m2:
            return m2.group(1).strip(), int(m2.group(2))
        return None
    name = m.group(1).strip()
    # Strip trailing seasonal note already captured partially
    name = re.sub(r"\s*-\s*High Season.*$", "", name, flags=re.I)
    name = re.sub(r"\s*-\s*Normal Season.*$", "", name, flags=re.I)
    stars = int(m.group(2))
    if name.lower() in {"room price", "e/bed"}:
        return None
    return name, stars


def parse_hotels_and_terms():
    wb = load_workbook(HOTEL_XLSX, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    city = None
    hotels = []
    current = None
    season = None  # High Season | Normal Season | None
    terms_lines: list[str] = []
    in_terms = False

    city_map = {
        "KUALA LUMPUR": "Kuala Lumpur",
        "GENTING": "Genting Highlands",
        "LANGKAWI": "Langkawi",
    }

    for row in rows:
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        a_s = str(a).strip() if a not in (None, "") else ""
        b_s = str(b).strip() if b not in (None, "") else ""

        if a_s.upper() in city_map:
            city = city_map[a_s.upper()]
            current = None
            season = None
            continue

        if a_s.lower().startswith("terms"):
            in_terms = True
            current = None
            continue

        if in_terms:
            line = a_s or b_s
            if line:
                terms_lines.append(line)
            continue

        if not city:
            continue

        # Hotel header row: name in A, "Room Price" in B
        if a_s and (b_s.lower() == "room price" or str(b).lower() == "room price"):
            hdr = parse_hotel_header(a_s)
            if not hdr:
                continue
            name, stars = hdr
            season = None
            low = a_s.lower()
            if "high season" in low:
                season = "High Season"
            elif "normal season" in low:
                season = "Normal Season"
            # Find or create hotel
            current = next((h for h in hotels if h["name"].lower() == name.lower() and h["city"] == city), None)
            if not current:
                current = {
                    "name": name,
                    "city": city,
                    "country": "Malaysia",
                    "starCategory": stars,
                    "rooms": [],
                }
                hotels.append(current)
            continue

        # Room row: name + numeric prices
        if current and a_s and isinstance(b, (int, float)):
            usd = round_usd(b)
            ebed = round_usd(c) if isinstance(c, (int, float)) else None
            room_name = a_s
            if season:
                room_name = f"{room_name} — {season}"
            # Dash high/normal season windows
            valid_from = KTH_VALID_FROM
            valid_to = HOTEL_TERMS_VALID_TO
            windows = None
            if season == "High Season":
                windows = [
                    {"validFrom": "2026-06-01", "validTo": "2026-06-30"},
                    {"validFrom": "2026-12-01", "validTo": "2026-12-20"},
                ]
            elif season == "Normal Season":
                # Rest of dates within hotel validity — use full window; picker uses room variant
                windows = [{"validFrom": KTH_VALID_FROM, "validTo": HOTEL_TERMS_VALID_TO}]
            current["rooms"].append(
                {
                    "name": room_name,
                    "usd": usd,
                    "extraBedUsd": ebed,
                    "inr": inr_from_usd(usd),
                    "extraBedInr": inr_from_usd(ebed) if ebed is not None else None,
                    "season": season,
                    "validFrom": valid_from,
                    "validTo": valid_to,
                    "windows": windows,
                }
            )

    hotel_terms = "\n".join(terms_lines).strip()
    return hotels, hotel_terms


def is_terms_row(name: str) -> bool:
    n = (name or "").strip().lower()
    return n.startswith("terms") or n.startswith("1.") or n.startswith("2.") or n.startswith("*")


def parse_kth_transfers(sheet_name: str, city: str, col_map: dict):
    wb = load_workbook(KTH_XLSX, data_only=True)
    ws = wb[sheet_name]
    transfers = []
    terms: list[str] = []
    in_terms = False
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        # Name usually in column B (index 1)
        name = cells[1] if len(cells) > 1 else None
        name_s = str(name).strip() if name not in (None, "") else ""
        if not name_s:
            continue
        if "terms" in name_s.lower() and "condition" in name_s.lower():
            in_terms = True
            continue
        if in_terms:
            if name_s:
                terms.append(name_s)
            continue
        if name_s.upper().startswith("DESTINATION") or "RATE SHEET" in name_s.upper():
            continue
        if name_s.lower().startswith("valid from"):
            continue
        if name_s.upper() in {"NAME", "TRANSFER FROM KLIA/KLIA 2:", "OVERLAND TRANSFER FROM PENANG"}:
            continue

        car = cells[col_map["car"]] if col_map["car"] < len(cells) else None
        if not isinstance(car, (int, float)):
            continue
        if is_terms_row(name_s):
            continue

        def g(key):
            idx = col_map.get(key)
            if idx is None or idx >= len(cells):
                return None
            v = cells[idx]
            return float(v) if isinstance(v, (int, float)) else None

        myr_car = float(car)
        myr_van10 = g("van10")
        myr_van18 = g("van18")
        myr_guide = g("van18Guide")
        transfers.append(
            {
                "name": re.sub(r"\s+", " ", name_s).strip(),
                "city": city,
                "country": "Malaysia",
                "myr": {
                    "car": myr_car,
                    "van10": myr_van10,
                    "van18": myr_van18,
                    "van18Guide": myr_guide,
                },
                "inr": {
                    "car": inr_from_myr(myr_car),
                    "van10": inr_from_myr(myr_van10),
                    "van18": inr_from_myr(myr_van18),
                    "van18Guide": inr_from_myr(myr_guide),
                },
            }
        )
    return transfers, "\n".join(terms).strip()


def parse_tickets_guides():
    wb = load_workbook(KTH_XLSX, data_only=True)
    tickets = []
    guides = []
    ticket_terms = []
    guide_terms = []

    if "TICKET" in wb.sheetnames:
        ws = wb["TICKET"]
        in_terms = False
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            name = None
            for cell in cells:
                if isinstance(cell, str) and cell.strip() and not cell.strip().replace(".", "", 1).isdigit():
                    # prefer longer descriptive
                    if name is None or len(str(cell)) > len(str(name)):
                        name = cell
            # Find adult/child numbers - scan numeric cells
            nums = [float(c) for c in cells if isinstance(c, (int, float))]
            name_s = str(name).strip() if name else ""
            if not name_s:
                continue
            if "terms" in name_s.lower():
                in_terms = True
                continue
            if in_terms:
                ticket_terms.append(name_s)
                continue
            if len(nums) < 1:
                continue
            if name_s.lower().startswith("valid") or "rate sheet" in name_s.lower():
                continue
            adult = nums[0]
            child = nums[1] if len(nums) > 1 else None
            tickets.append(
                {
                    "name": re.sub(r"\s+", " ", name_s),
                    "city": "Kuala Lumpur",
                    "country": "Malaysia",
                    "adultMyr": adult,
                    "childMyr": child,
                    "adultInr": inr_from_myr(adult),
                    "childInr": inr_from_myr(child),
                }
            )

    if "GUIDE" in wb.sheetnames:
        ws = wb["GUIDE"]
        in_terms = False
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            name = cells[1] if len(cells) > 1 else cells[0] if cells else None
            name_s = str(name).strip() if name not in (None, "") else ""
            nums = [float(c) for c in cells if isinstance(c, (int, float))]
            if not name_s:
                continue
            if "terms" in name_s.lower():
                in_terms = True
                continue
            if in_terms:
                guide_terms.append(name_s)
                continue
            if not nums:
                continue
            if name_s.lower().startswith("valid") or "rate" in name_s.lower() and "sheet" in name_s.lower():
                continue
            myr = nums[0]
            guides.append(
                {
                    "name": re.sub(r"\s+", " ", name_s),
                    "city": "Kuala Lumpur",
                    "country": "Malaysia",
                    "myr": myr,
                    "inr": inr_from_myr(myr),
                }
            )

    return tickets, guides, "\n".join(ticket_terms), "\n".join(guide_terms)


def main():
    hotels, hotel_terms = parse_hotels_and_terms()

    # KTH sheet columns differ by city
    kl_t, kl_terms = parse_kth_transfers(
        "KUALA LUMPUR ",
        "Kuala Lumpur",
        {"car": 2, "van10": 3, "van18": 4, "van18Guide": 5},
    )
    lg_t, lg_terms = parse_kth_transfers(
        "LANGKAWI",
        "Langkawi",
        {"car": 3, "van10": 4, "van18": 5, "van18Guide": 6},
    )
    pg_t, pg_terms = parse_kth_transfers(
        "PENANG",
        "Penang",
        {"car": 3, "van10": 4, "van18": 5, "van18Guide": 6},
    )
    tickets, guides, ticket_terms, guide_terms = parse_tickets_guides()

    # Filter accidental terms-as-product
    transfers = [t for t in (kl_t + lg_t + pg_t) if "terms" not in t["name"].lower()]

    out = {
        "source": {
            "hotels": "Malaysia Contracted rates .xlsx",
            "kth": "KTH RATE SHEET 2026 - TREVIO.xlsx",
            "hotelValidFrom": KTH_VALID_FROM,
            "hotelValidTo": HOTEL_TERMS_VALID_TO,
            "kthValidFrom": KTH_VALID_FROM,
            "kthValidTo": KTH_VALID_TO,
            "fx": {"USD_TO_INR": USD_TO_INR, "MYR_TO_INR": MYR_TO_INR},
            "roeNote": "Rate of exchange considered as on 23rd March 26 (per hotel sheet).",
        },
        "terms": {
            "hotels": hotel_terms,
            "kthKualaLumpur": kl_terms,
            "kthLangkawi": lg_terms,
            "kthPenang": pg_terms,
            "tickets": ticket_terms,
            "guides": guide_terms,
        },
        "hotels": hotels,
        "transfers": transfers,
        "tickets": tickets,
        "guides": guides,
    }
    OUT.write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")
    print(
        json.dumps(
            {
                "out": str(OUT),
                "hotels": len(hotels),
                "rooms": sum(len(h["rooms"]) for h in hotels),
                "transfers": len(transfers),
                "tickets": len(tickets),
                "guides": len(guides),
                "sample": hotels[0] if hotels else None,
                "dash": next((h for h in hotels if "Dash" in h["name"]), None),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
